from pathlib import Path
import json, re
import openpyxl

class EffectLibrary:
    def __init__(self, xlsx_path):
        self.xlsx_path = Path(xlsx_path)
        self.rows=[]
        wb=openpyxl.load_workbook(self.xlsx_path, read_only=True, data_only=True)
        for sheet in wb.sheetnames:
            ws=wb[sheet]
            headers=[c.value for c in next(ws.iter_rows())]
            for vals in ws.iter_rows(min_row=2, values_only=True):
                d=dict(zip(headers, vals)); d['sheet']=sheet
                self.rows.append(d)
        self.by_shortcut={r['Shortcut']:r for r in self.rows}
    def lookup(self, terms, limit=5):
        terms=[t.lower() for t in terms]
        scored=[]
        for r in self.rows:
            text=' '.join(str(r.get(k,'')) for k in ['Shortcut','Capability','Primary Use / Intent','Scene Recipe','Visual / Execution Notes','Camera / Scale','Ad Role']).lower()
            score=sum(1 for t in terms if t in text)
            if score: scored.append((score,r))
        scored.sort(key=lambda x:(-x[0], x[1]['Shortcut']))
        return [r for _,r in scored[:limit]]
    def get(self, shortcut): return self.by_shortcut.get(shortcut)
    def export_json(self, path):
        Path(path).write_text(json.dumps(self.rows, ensure_ascii=False, indent=2), encoding='utf-8')
